"""
Review report: one row per finding. This spreadsheet is the human
checkpoint (see ADR-006 for why it's .xlsx and not .csv).

Columns cover both the flag itself and enough disambiguating context
(dates, places, parents, spouses, and cross-reference IDs) to make it
easy to find the *correct* person in your genealogy program when
several relatives share a name -- which is common across generations.

Identifier columns are populated only when present in the source
GEDCOM. Nothing is inferred or guessed.

Columns (roughly in review order -- what to decide on first, on the
left; disambiguation/identifier lookups you only need occasionally,
further right):
    action             - YOU FILL THIS IN: keep | redact | rewrite:<text> |
        mask | privatize. Column A, frozen alongside the header row, so it
        stays visible while you scroll right to read context. Has an in-cell
        dropdown listing the five options (pick "rewrite:" then type your
        own text after it for a custom rewrite).
    name               - whose record the text is in: a person, "A & B"
        for a family, the people sharing a shared note, a source title
    found_in           - where in the record, e.g. "Death > Cause"
    severity, category, snippet, matched_term - what got flagged and why;
        snippet is the primary thing to read to make a keep/redact call,
        and wraps within its cell instead of running off-screen. The
        matched term is highlighted (bold + colored) within the snippet
        as rich text -- see ADR-006. Severity shows as a red/yellow/green
        circle icon; hover the cell for the actual word as a tooltip.
    sex, birth, death  - vital info to disambiguate same-named relatives
    parents, spouses   - resolved from FAM records, same reason
    familysearch_id / familysearch_url
    findagrave_id / findagrave_url
    ancestral_file_number (AFN), reference_number (REFN)
    xref               - GEDCOM pointer, e.g. @I42@ -- not human-meaningful,
        kept at the tail with the other rarely-needed lookup/bookkeeping
        columns
    line, line_end     - source line numbers, counting from 1 (needed by
        the redactor)

A second sheet, "Summary", says what was scanned (and what wasn't), the
preset and rules used, counts per category, and any warnings from
reading the file.

Everything written comes from the GEDCOM, which may be hostile: no
value is ever written as a formula (see ``_force_text``), and URLs are built
only from identifiers that match the expected format.
"""

from __future__ import annotations

import os
import re
from collections import Counter

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .parser import Gedcom, Individual
from .parser import TextBlock as GedTextBlock
from .scanner import Finding

FAMILYSEARCH_PERSON_URL = "https://www.familysearch.org/tree/person/details/{id}"
FINDAGRAVE_URL = "https://www.findagrave.com/memorial/{id}"
_FAMILYSEARCH_ID = re.compile(r"^[A-Z0-9]{4}-[A-Z0-9]{2,4}$")
_FINDAGRAVE_ID = re.compile(r"^[0-9]{1,12}$")

# Bold, dark red -- applied to just the matched term within the
# snippet cell so it's visually distinct from its surrounding context.
HIGHLIGHT_FONT = InlineFont(b=True, color="FFC00000")

# Snippet cells wrap instead of running off-screen as one long line.
_WRAP_ALIGNMENT = Alignment(wrap_text=True, vertical="top")
_CENTER_ALIGNMENT = Alignment(horizontal="center", vertical="top")
_DATA_ROW_HEIGHT = 60  # tall enough for a wrapped ~200-char snippet

# Severity shows as a compact icon in the cell; the actual word is a
# hover tooltip (Excel cell comment) since the icon alone isn't
# self-explanatory.
SEVERITY_ICONS = {
    "high": "\U0001f534",
    "medium": "\U0001f7e1",
    "low": "\U0001f7e2",
}  # red/yellow/green circle

# Dropdown offered on the "action" column. "rewrite:" is included as a
# starting point -- picking it, then typing your own text after it, is
# how you get "rewrite:<text>"; the dropdown doesn't block free typing
# (no error message/style set), so a full custom rewrite: value is
# still fine even though it won't exactly match a list entry.
_ACTION_OPTIONS = '"keep,redact,mask,privatize,rewrite:"'

# Columns that hold integer line numbers rather than text, so they get
# written as real numbers (sortable/filterable) instead of strings.
_INT_FIELDS = {"line", "line_end"}

# A cell starting with one of these can be read as a formula by Excel
# or by whatever the sheet is later exported to.
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")

FIELDNAMES = [
    "action",
    "name",
    "found_in",
    "severity",
    "category",
    "snippet",
    "matched_term",
    "sex",
    "birth",
    "death",
    "parents",
    "spouses",
    "familysearch_id",
    "familysearch_url",
    "findagrave_id",
    "findagrave_url",
    "ancestral_file_number",
    "reference_number",
    "xref",
    "line",
    "line_end",
]

SEVERITY_ORDER = {"high": 0, "medium": 1, "low": 2}

# Plain words for the tags most often seen in a "found_in" path.
_TAG_WORDS = {
    "NOTE": "Note",
    "CAUS": "Cause",
    "TEXT": "Text",
    "PAGE": "Citation detail",
    "FILE": "Media file",
    "SOUR": "Source citation",
    "OBJE": "Media",
    "BIRT": "Birth",
    "CHR": "Christening",
    "BAPM": "Baptism",
    "DEAT": "Death",
    "BURI": "Burial",
    "CREM": "Cremation",
    "MARR": "Marriage",
    "DIV": "Divorce",
    "RESI": "Residence",
    "OCCU": "Occupation",
    "EDUC": "Education",
    "MILI": "Military",
    "_MILT": "Military",
    "IMMI": "Immigration",
    "EMIG": "Emigration",
    "NATU": "Naturalization",
    "CENS": "Census",
    "PROB": "Probate",
    "WILL": "Will",
    "EVEN": "Event",
    "FACT": "Fact",
    "DATA": "Data",
    "ADDR": "Address",
    "PHON": "Phone",
    "EMAIL": "Email",
    "NAME": "Name",
}
_RECORD_WORDS = {
    "INDI": "Person",
    "FAM": "Family",
    "NOTE": "Shared note",
    "SOUR": "Source",
    "REPO": "Repository",
    "OBJE": "Media",
    "SUBM": "Submitter",
    "HEAD": "File header",
}

# What the scanner reads, in the order the summary lists it.
LOCATION_LABELS = {
    "note": "Notes: on people, families, events, sources, and shared NOTE records",
    "cause": "Cause of death or of an event (CAUS)",
    "source_text": "Text quoted from a source (TEXT)",
    "citation": "Citation details (PAGE)",
    "media_file": "Media file paths (FILE)",
    "submitter": "Submitter record: your name and contact details (SUBM)",
    "address": "Addresses, phones, emails and web sites",
    "custom_tag": "Program-specific fields (_TAG)",
}
NOT_SCANNED = (
    "Names, dates, places, occupations and other short fields are not scanned, and are "
    "copied to the cleaned file unchanged. Skim them yourself, and remove living people "
    "in your genealogy program, which already does that."
)


def _vital(date: str, place: str) -> str:
    if date and place:
        return f"{date} in {place}"
    return date or place or ""


def familysearch_url(fsid: str) -> str:
    fsid = fsid.strip().upper()
    return FAMILYSEARCH_PERSON_URL.format(id=fsid) if _FAMILYSEARCH_ID.match(fsid) else ""


def findagrave_url(memorial: str) -> str:
    memorial = memorial.strip()
    return FINDAGRAVE_URL.format(id=memorial) if _FINDAGRAVE_ID.match(memorial) else ""


def _person_context(person: Individual) -> dict:
    return {
        "xref": person.xref,
        "name": person.name,
        "sex": person.sex,
        "birth": _vital(person.birth_date, person.birth_place),
        "death": _vital(person.death_date, person.death_place),
        "parents": "; ".join(person.parent_names),
        "spouses": "; ".join(person.spouse_names),
        "familysearch_id": person.fsftid,
        "familysearch_url": familysearch_url(person.fsftid),
        "findagrave_id": person.findagrave,
        "findagrave_url": findagrave_url(person.findagrave),
        "ancestral_file_number": person.afn,
        "reference_number": person.refn,
    }


def _names(xrefs: list[str], ged: Gedcom) -> str:
    people = ged.by_xref
    fams = ged.families
    names: list[str] = []
    for x in dict.fromkeys(xrefs):
        if x in people:
            names.append(people[x].name)
        elif x in fams:
            names.append(_family_name(fams[x].husb, fams[x].wife, ged))
        elif x in ged.record_names:
            names.append(ged.record_names[x])
    if len(names) > 3:
        return "; ".join(names[:3]) + f"; and {len(names) - 3} more"
    return "; ".join(names)


def _family_name(husb: str, wife: str, ged: Gedcom) -> str:
    people = ged.by_xref
    parts = [people[x].name for x in (husb, wife) if x in people]
    return " & ".join(parts) or "(family)"


def _context(block: GedTextBlock, ged: Gedcom) -> dict:
    """Who and what the text belongs to."""
    people = ged.by_xref
    rec = block.record_xref
    if block.record_type == "INDI" and rec in people:
        return _person_context(people[rec])
    if block.record_type == "NOTE":
        refs = ged.note_refs.get(rec, [])
        if len(set(refs)) == 1 and refs[0] in people:
            return {**_person_context(people[refs[0]]), "xref": rec}
        return {"xref": rec, "name": _names(refs, ged) or "(not linked to anyone)"}
    if block.record_type == "FAM" and rec in ged.families:
        fam = ged.families[rec]
        return {"xref": rec, "name": _family_name(fam.husb, fam.wife, ged)}
    if rec in ged.record_names:
        return {"xref": rec, "name": ged.record_names[rec]}
    return {"xref": rec, "name": _RECORD_WORDS.get(block.record_type, block.record_type)}


def found_in(block: GedTextBlock) -> str:
    record = _RECORD_WORDS.get(block.record_type, block.record_type.title())
    if block.record_type == "NOTE" and block.level == 0:
        return record
    steps = [_TAG_WORDS.get(t, t) for t in block.path]
    return " > ".join([record, *steps])


def _rich_snippet(snippet: str, matched_term: str, at: int = -1) -> CellRichText | str:
    """Return the snippet with matched_term highlighted as a bold/colored
    rich-text run, or plain snippet text if the term can't be located
    (shouldn't happen -- snippet is always built from the same source
    text as matched_term -- but fail safe rather than raise)."""
    term_len = len(matched_term)
    if not (0 <= at and snippet[at : at + term_len].lower() == matched_term.lower()):
        at = snippet.lower().find(matched_term.lower())
    if at == -1 or not matched_term:
        return snippet
    before, hit, after = snippet[:at], snippet[at : at + term_len], snippet[at + term_len :]
    parts = [p for p in (before,) if p] + [TextBlock(HIGHLIGHT_FONT, hit)]
    parts += [p for p in (after,) if p]
    return CellRichText(*parts)


def _write_rows(ws: Worksheet, rows: list[dict]) -> None:
    ws.append(FIELDNAMES)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "B2"  # keep the "action" column and header visible while scrolling
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(FIELDNAMES)).coordinate}"

    for row in rows:
        values = []
        for field in FIELDNAMES:
            if field == "snippet":
                values.append(_rich_snippet(row["snippet"], row["matched_term"], row["match_at"]))
            elif field == "severity":
                values.append(SEVERITY_ICONS.get(row["severity"], row["severity"]))
            elif field in _INT_FIELDS:
                values.append(int(row[field]))
            else:
                values.append(row.get(field, ""))
        ws.append(values)

        out_row = ws[ws.max_row]
        for cell in out_row:
            _force_text(cell)
        out_row[FIELDNAMES.index("snippet")].alignment = _WRAP_ALIGNMENT
        severity_cell = out_row[FIELDNAMES.index("severity")]
        severity_cell.alignment = _CENTER_ALIGNMENT
        severity_cell.comment = Comment(row["severity"].capitalize(), "GEDCOM Privacy Guard")
        ws.row_dimensions[ws.max_row].height = _DATA_ROW_HEIGHT

    if rows:
        dv = DataValidation(
            type="list", formula1=_ACTION_OPTIONS, allow_blank=True, showDropDown=False
        )
        dv.prompt = "Choose an action, or type rewrite:<your text>"
        dv.promptTitle = "Action"
        dv.showInputMessage = True
        ws.add_data_validation(dv)
        dv.add(f"A2:A{ws.max_row}")

    widths = {
        "snippet": 80,
        "name": 20,
        "found_in": 18,
        "parents": 24,
        "spouses": 24,
        "matched_term": 16,
        "severity": 10,
        "action": 20,
    }
    for i, field in enumerate(FIELDNAMES, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(field, 14)


def _force_text(cell) -> None:
    """openpyxl turns any string starting with '=' into a formula. Store
    such values as plain text, and mark every risky value with Excel's
    quote prefix so editing the cell later doesn't turn it into one."""
    value = cell.value
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        cell.data_type = "s"
        cell.quotePrefix = True


def build_rows(ged: Gedcom, scanners: list, progress=None) -> tuple[list[dict], Counter]:
    """One row per finding, most severe first, and how many text fields
    were scanned in each location. ``progress(fraction)``, if given, is
    called about a hundred times as the scan goes."""
    rows = []
    scanned: Counter = Counter()
    total = len(ged.text_blocks)
    step = max(1, total // 100)
    for i, block in enumerate(ged.text_blocks):
        if progress and i % step == 0:
            progress(i / total)
        scanned[block.location] += 1
        findings: list[Finding] = []
        for scanner in scanners:
            findings.extend(scanner.scan(block.text, block.location))
        if not findings:
            continue
        context = _context(block, ged)
        where = found_in(block)
        for finding in findings:
            rows.append(
                {
                    **context,
                    "found_in": where,
                    "line": block.start_line + 1,
                    "line_end": block.end_line + 1,
                    "category": finding.category,
                    "severity": finding.severity,
                    "matched_term": finding.matched_term,
                    "snippet": finding.snippet,
                    "match_at": finding.match_at,
                    "action": "",  # left blank for human review
                }
            )
    rows.sort(key=lambda r: (SEVERITY_ORDER.get(r["severity"], 9), r["name"], r["line"]))
    return rows, scanned


def _write_summary(ws: Worksheet, ged: Gedcom, gedcom_path: str, rows, scanned, scanners):
    bold = Font(bold=True)

    def heading(text):
        ws.append([])
        ws.append([text])
        ws[ws.max_row][0].font = bold

    ws.append(["GEDCOM Privacy Guard review"])
    ws[1][0].font = Font(bold=True, size=14)
    ws.append(["File", os.path.basename(gedcom_path)])
    ws.append(["Read as", ged.encoding])
    ws.append(["People", len(ged.individuals)])
    ws.append(["Findings", len(rows)])
    rules = getattr(scanners[0], "rules", None) if scanners else None
    if rules is not None:
        ws.append(["Preset", rules.preset])
        ws.append(["Your rules file", os.path.basename(rules.source) or "(none: built-in rules)"])

    heading("What was scanned")
    ws.append(["Where", "Text fields scanned"])
    for key, label in LOCATION_LABELS.items():
        ws.append([label, scanned.get(key, 0)])
    ws.append([NOT_SCANNED])

    heading("Findings by category")
    counts = Counter(r["category"] for r in rows)
    if rules is not None:
        for cat in rules.categories:
            state = counts.get(cat.name, 0) if cat.enabled else "off"
            ws.append([cat.name, state, cat.description])
    else:
        for name, n in counts.most_common():
            ws.append([name, n])

    if ged.warnings:
        heading("Warnings from reading the file")
        for w in ged.warnings:
            ws.append([w])

    for row in ws.iter_rows():
        for cell in row:
            _force_text(cell)
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 80


def write_report(path: str, ged: Gedcom, scanners: list, gedcom_path: str = "", built=None) -> int:
    """Write review.xlsx and return the number of findings. ``built`` is
    build_rows' result, if the caller already has it."""
    rows, scanned = built or build_rows(ged, scanners)

    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    _write_rows(ws, rows)
    _write_summary(wb.create_sheet("Summary"), ged, gedcom_path, rows, scanned, scanners)
    wb.save(path)

    return len(rows)


def read_reviewed_report(path: str) -> list[dict]:
    """Rows of a reviewed report. Reports from before 0.4 used 0-based
    note_line_start/note_line_end; they're converted to line/line_end."""
    wb = load_workbook(path, data_only=True)
    try:
        ws = wb["Review"] if "Review" in wb.sheetnames else wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = list(next(rows_iter, []))
        rows = [
            dict(zip(header, row, strict=False))
            for row in rows_iter
            if any(v is not None for v in row)
        ]
    finally:
        wb.close()
    for row in rows:
        if "line" not in row and "note_line_start" in row:
            row["line"] = int(row["note_line_start"]) + 1
            row["line_end"] = int(row["note_line_end"]) + 1
    return rows

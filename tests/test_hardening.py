"""A GEDCOM is untrusted input, and a user's files are precious."""

import pytest
from conftest import SAMPLE, ged
from openpyxl import load_workbook

from gedcom_privacy_guard.cli import main
from gedcom_privacy_guard.parser import read_gedcom
from gedcom_privacy_guard.report import FIELDNAMES, familysearch_url, findagrave_url, write_report
from gedcom_privacy_guard.scanner import KeywordScanner

# ---- spreadsheet formula injection ----

HOSTILE = ged(
    "0 @I1@ INDI",
    '1 NAME =HYPERLINK("http://evil.example","click") /x/',
    "1 SEX +cmd|' /C calc'!A0",
    "1 BIRT",
    "2 DATE @SUM(1+1)",
    "2 PLAC -2+3",
    "1 _FSFTID =1+1",
    "1 NOTE =1+1 he was arrested",
)


def test_values_are_never_formulas(write_ged, tmp_path):
    out = str(tmp_path / "review.xlsx")
    write_report(out, read_gedcom(write_ged(HOSTILE)), [KeywordScanner()])
    wb = load_workbook(out)
    ws = wb["Review"]
    row = {name: ws.cell(row=2, column=i + 1) for i, name in enumerate(FIELDNAMES)}
    for field in ("name", "sex", "birth", "familysearch_id"):
        cell = row[field]
        assert cell.data_type == "s", field
        assert cell.quotePrefix, field
    assert row["name"].value.startswith("=HYPERLINK")  # shown as text, unchanged
    for sheet in wb.worksheets:
        for r in sheet.iter_rows():
            assert all(c.data_type != "f" for c in r)
    wb.close()


# ---- identifiers only become URLs when they look right ----


@pytest.mark.parametrize(
    ("fsid", "ok"),
    [("LZXY-ABC", True), ("kwj1-23q", True), ("LZXY-ABC/../x", False), ("javascript:x", False)],
)
def test_familysearch_url_validation(fsid, ok):
    assert bool(familysearch_url(fsid)) is ok


@pytest.mark.parametrize(("memorial", "ok"), [("12345", True), ("123?x=1", False), ("", False)])
def test_findagrave_url_validation(memorial, ok):
    assert bool(findagrave_url(memorial)) is ok


# ---- the CLI never overwrites what you gave it ----


def _copy_sample(tmp_path):
    path = tmp_path / "tree.ged"
    path.write_bytes(open(SAMPLE, "rb").read())
    return str(path)


def test_scan_refuses_to_overwrite_the_gedcom(tmp_path, capsys):
    path = _copy_sample(tmp_path)
    before = open(path, "rb").read()
    assert main(["scan", path, "--out", path, "--force"]) == 1
    assert "is your GEDCOM file" in capsys.readouterr().err
    assert open(path, "rb").read() == before


def test_scan_wont_replace_a_reviewed_report_without_force(tmp_path, capsys):
    path = _copy_sample(tmp_path)
    report = tmp_path / "review.xlsx"
    report.write_bytes(b"my decisions")
    assert main(["scan", path, "--out", str(report)]) == 1
    assert "--force" in capsys.readouterr().err
    assert report.read_bytes() == b"my decisions"
    assert main(["scan", path, "--out", str(report), "--force"]) == 0


@pytest.mark.parametrize("target", ["gedcom", "report"])
def test_redact_refuses_to_overwrite_its_inputs(tmp_path, capsys, target):
    path = _copy_sample(tmp_path)
    report = str(tmp_path / "review.xlsx")
    assert main(["scan", path, "--out", report]) == 0
    out = path if target == "gedcom" else report
    before = open(out, "rb").read()
    assert main(["redact", path, "--report", report, "--out", out]) == 1
    assert "Choose a different output file" in capsys.readouterr().err
    assert open(out, "rb").read() == before


def test_full_cli_round_trip(tmp_path, capsys):
    path = _copy_sample(tmp_path)
    report = str(tmp_path / "review.xlsx")
    out = str(tmp_path / "cleaned.ged")
    assert main(["scan", path, "--out", report]) == 0
    wb = load_workbook(report)
    ws = wb["Review"]
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1, value="redact")
    wb.save(report)
    assert main(["redact", path, "--report", report, "--out", out]) == 0
    cleaned = open(out, encoding="utf-8").read()
    assert "alcoholism" not in cleaned and "anxiety disorder" not in cleaned
    assert "Original .ged was not modified" in capsys.readouterr().out


# ---- expected failures are one plain line, never a traceback ----


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (b"", "is empty"),
        (b"   \n\n", "is empty"),
        (b"PK\x03\x04 a zip file", "doesn't look like a GEDCOM"),
        (b"Name,Born\nJane,1900\n", "doesn't look like a GEDCOM"),
    ],
)
def test_not_a_gedcom(tmp_path, capsys, content, message):
    path = tmp_path / "x.ged"
    path.write_bytes(content)
    assert main(["scan", str(path), "--out", str(tmp_path / "r.xlsx")]) == 1
    err = capsys.readouterr().err
    assert message in err
    assert "Traceback" not in err


def test_missing_file(tmp_path, capsys):
    assert main(["scan", str(tmp_path / "nope.ged"), "--out", str(tmp_path / "r.xlsx")]) == 1
    assert "file not found" in capsys.readouterr().err


def test_report_that_isnt_a_spreadsheet(tmp_path, capsys):
    path = _copy_sample(tmp_path)
    bad = tmp_path / "review.xlsx"
    bad.write_text("not a spreadsheet")
    assert main(["redact", path, "--report", str(bad), "--out", str(tmp_path / "o.ged")]) == 1
    assert "Can't read" in capsys.readouterr().err


def test_bad_rules_file(tmp_path, capsys):
    rules = tmp_path / "r.yaml"
    rules.write_text("extends: strict\n")
    assert main(["rules", "--rules", str(rules)]) == 1
    assert "must be open, standard, cautious or none" in capsys.readouterr().err


def test_rules_command_lists_categories(capsys):
    assert main(["rules", "--preset", "cautious", "-v"]) == 0
    out = capsys.readouterr().out
    assert "[on ] family_secrets" in out and "Rules are valid." in out

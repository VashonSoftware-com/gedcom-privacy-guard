"""End-to-end behavior: parse -> scan -> report -> redact."""

import pytest
from conftest import SAMPLE, ged
from openpyxl import load_workbook

from gedcom_privacy_guard.parser import read_gedcom
from gedcom_privacy_guard.redactor import (
    REDACTED_TEXT,
    InvalidActionError,
    ReportMismatchError,
    apply_redactions,
)
from gedcom_privacy_guard.report import FIELDNAMES, read_reviewed_report, write_report
from gedcom_privacy_guard.rules import load_rules
from gedcom_privacy_guard.scanner import KeywordScanner


def _scanner(preset="standard"):
    return KeywordScanner(load_rules(preset=preset))


def _rows(path, scanners=None, tmp_path=None):
    """Scan a GEDCOM into a report and read the rows back."""
    out = str(tmp_path / "review.xlsx")
    write_report(out, read_gedcom(path), scanners or [_scanner()])
    return read_reviewed_report(out)


def _row(line, line_end=None, action="redact", term="", name="Jane Doe"):
    return {
        "line": line,
        "line_end": line_end or line,
        "action": action,
        "matched_term": term,
        "name": name,
    }


# ---- parsing ----


def test_parse_finds_four_individuals():
    individuals = read_gedcom(SAMPLE).individuals
    assert len(individuals) == 4
    names = [p.name for p in individuals]
    assert names.count("Robert Smith") == 2  # two same-named, different people


def test_disambiguates_same_named_individuals_by_vitals():
    roberts = [p for p in read_gedcom(SAMPLE).individuals if p.name == "Robert Smith"]
    assert len({p.birth_date for p in roberts}) == 2  # distinguishable by birth date


def test_resolves_parents_and_spouses_from_fam_records():
    by_xref = read_gedcom(SAMPLE).by_xref
    assert set(by_xref["@I3@"].parent_names) == {"John Smith", "Mary Callahan"}
    assert by_xref["@I1@"].spouse_names == ["Mary Callahan"]


def test_captures_familysearch_id_when_present():
    by_xref = read_gedcom(SAMPLE).by_xref
    assert by_xref["@I1@"].fsftid == "LZXY-ABC"
    assert by_xref["@I2@"].fsftid == ""  # not linked -- must stay blank, not guessed


def test_parse_captures_multiline_note():
    john = next(p for p in read_gedcom(SAMPLE).individuals if p.name == "John Smith")
    assert len(john.notes) == 1
    assert "alcoholism" in john.notes[0].text
    assert "institutionalized" in john.notes[0].text


def test_unhandled_level1_tag_does_not_corrupt_death_date(write_ged):
    # CHAN (record last-modified date) appears in almost every real
    # GEDCOM export. Its DATE child must never leak into death_date.
    path = write_ged(
        ged(
            "0 @I1@ INDI",
            "1 NAME Jane /Doe/",
            "1 BIRT",
            "2 DATE 01 JAN 1900",
            "1 DEAT",
            "2 DATE 02 FEB 1970",
            "1 CHAN",
            "2 DATE 15 MAR 2020",
        )
    )
    person = read_gedcom(path).individuals[0]
    assert person.death_date == "02 FEB 1970"
    assert person.birth_date == "01 JAN 1900"


def test_conc_keeps_spaces_at_line_ends(write_ged):
    path = write_ged(ged("0 @I1@ INDI", "1 NOTE He was a ", "2 CONC drunkard."))
    assert read_gedcom(path).individuals[0].notes[0].text == "He was a drunkard."


# ---- every text location is scanned ----

EVERYWHERE = ged(
    "0 @I1@ INDI",
    "1 NAME Jane /Doe/",
    "1 DEAT",
    "2 CAUS suicide",
    "2 NOTE she had depression",
    "1 NOTE @N1@",
    "1 SOUR @S1@",
    "2 PAGE SSN 123-45-6789",
    "2 DATA",
    "3 TEXT arrested in 1950",
    "1 OBJE",
    "2 FILE C:\\Users\\jdoe\\Pictures\\jane.jpg",
    "1 _MYNOTE was an addict",
    "1 FAMS @F1@",
    "0 @I2@ INDI",
    "1 NAME John /Doe/",
    "1 FAMS @F1@",
    "0 @F1@ FAM",
    "1 HUSB @I2@",
    "1 WIFE @I1@",
    "1 MARR",
    "2 NOTE married after the affair",
    "0 @N1@ NOTE a rumor says",
    "1 CONC  she was in prison",
    "0 @S1@ SOUR",
    "1 TITL County records",
    "1 TEXT she was convicted",
    "0 @U1@ SUBM",
    "1 NAME Researcher Name",
    "1 PHON 206-555-0100",
)


def test_scans_every_text_location(write_ged, tmp_path):
    rows = _rows(write_ged(EVERYWHERE), tmp_path=tmp_path)
    found = {(r["found_in"], r["matched_term"]) for r in rows}
    assert ("Person > Death > Cause", "suicide") in found
    assert ("Person > Death > Note", "depression") in found
    assert ("Shared note", "rumor") in found
    assert ("Shared note", "prison") in found  # text on the CONC line
    assert ("Person > Source citation > Citation detail", "123-45-6789") in found
    assert ("Person > Source citation > Data > Text", "arrested") in found
    assert ("Person > Media > Media file", "C:\\") in found
    assert ("Person > _MYNOTE", "addict") in found
    assert ("Family > Marriage > Note", "affair") in found
    assert ("Source > Text", "convicted") in found
    assert ("Submitter > Phone", "206-555-0100") in found


def test_rows_name_the_people_involved(write_ged, tmp_path):
    rows = _rows(write_ged(EVERYWHERE), tmp_path=tmp_path)
    by_where = {r["found_in"]: r["name"] for r in rows}
    assert by_where["Family > Marriage > Note"] == "John Doe & Jane Doe"
    assert by_where["Shared note"] == "Jane Doe"  # the one person pointing at it
    assert by_where["Source > Text"] == "County records"


def test_head_record_is_not_scanned(write_ged, tmp_path):
    path = write_ged(
        "0 HEAD\n1 SOUR Prison Records Software\n2 CORP Addict Ltd\n3 PHON 206-555-0100\n0 TRLR\n"
    )
    assert _rows(path, tmp_path=tmp_path) == []


def test_pointer_values_are_not_text(write_ged):
    path = write_ged(ged("0 @I1@ INDI", "1 NOTE @N1@", "1 OBJE @O1@"))
    assert read_gedcom(path).text_blocks == []


# ---- scanner ----


def test_scanner_flags_sensitive_terms():
    categories = {f.category for f in _scanner().scan("He struggled with alcoholism.")}
    assert "substance_use" in categories


def test_scanner_ignores_clean_text():
    assert _scanner().scan("He worked as a machinist in Ohio for 30 years.") == []


def test_great_depression_is_not_flagged():
    assert _scanner().scan("They lost the farm in the Great Depression.") == []
    assert _scanner().scan("She suffered from depression.") != []


def test_whole_word_matching():
    s = _scanner()
    assert s.scan("Veterans Assn. of Ohio") == []  # 'ssn' inside a word
    assert [f.matched_term for f in s.scan("His SSN was lost")] == ["SSN"]
    assert [f.matched_term for f in s.scan("an alcoholic father")] == ["alcoholic"]  # prefix


def test_old_legacy_format_codes_do_not_hide_terms():
    # Older Legacy notes wrap formatted words in «b»…«/b»; « and » count as word boundaries.
    assert [f.matched_term for f in _scanner().scan("«b»alcoholism«/b»")] == ["alcoholism"]


def test_phrase_matches_across_line_break():
    assert [f.category for f in _scanner().scan("anxiety\ndisorder")] == ["mental_health"]


def test_pii_patterns():
    s = _scanner("open")
    terms = [f.matched_term for f in s.scan("SSN 123-45-6789, call (206) 555-0100, a@b.org")]
    assert "123-45-6789" in terms
    assert "(206) 555-0100" in terms
    assert "a@b.org" in terms
    assert s.scan("Born 1918-1920; census 1910") == []


def test_already_private_text_is_flagged():
    findings = _scanner("open").scan("Her [[first husband]] and {the baby}")
    assert [f.matched_term for f in findings] == ["[[first husband]]", "{the baby}"]
    assert {f.category for f in findings} == {"already_private"}


def test_snippet_marks_the_exact_match():
    text = "prison guard. " * 20 + "He went to prison."
    finding = _scanner().scan(text)[-1]
    assert finding.snippet[finding.match_at : finding.match_at + 6] == "prison"
    assert finding.snippet.startswith("...")


# ---- report ----


def test_write_report_highlights_matched_term_in_snippet(tmp_path):
    path = str(tmp_path / "review.xlsx")
    write_report(path, read_gedcom(SAMPLE), [_scanner()])
    wb = load_workbook(path, rich_text=True)
    ws = wb["Review"]
    assert [c.value for c in ws[1]] == FIELDNAMES
    snippet_col = FIELDNAMES.index("snippet")
    matched_col = FIELDNAMES.index("matched_term")
    for row in ws.iter_rows(min_row=2, values_only=True):
        rich = row[snippet_col]
        highlighted = "".join(p.text for p in rich if not isinstance(p, str))
        assert highlighted.lower() == row[matched_col].lower()
    wb.close()


def test_report_has_summary_sheet(tmp_path):
    path = str(tmp_path / "review.xlsx")
    write_report(path, read_gedcom(SAMPLE), [_scanner()], gedcom_path=SAMPLE)
    wb = load_workbook(path)
    text = " ".join(str(c) for row in wb["Summary"].iter_rows(values_only=True) for c in row)
    wb.close()
    assert "sample.ged" in text
    assert "What was scanned" in text
    assert "not scanned" in text
    assert "standard" in text


def test_read_reviewed_report_round_trips_through_xlsx(tmp_path):
    path = str(tmp_path / "review.xlsx")
    count = write_report(path, read_gedcom(SAMPLE), [_scanner()])
    reviewed = read_reviewed_report(path)
    assert len(reviewed) == count
    assert all(field in reviewed[0] for field in FIELDNAMES)
    assert all(row["action"] is None for row in reviewed)


def test_reads_old_reports_with_zero_based_line_numbers(tmp_path):
    from openpyxl import Workbook

    path = str(tmp_path / "old.xlsx")
    wb = Workbook()
    wb.active.append(["action", "name", "note_line_start", "note_line_end"])
    wb.active.append(["redact", "Jane", 16, 19])
    wb.save(path)
    row = read_reviewed_report(path)[0]
    assert (row["line"], row["line_end"]) == (17, 20)


# ---- redaction ----

NOTE_GED = ged(
    "0 @I1@ INDI",
    "1 NAME Jane /Doe/",
    "1 NOTE He struggled with alcoholism and a drug problem too.",
)
NOTE_LINE = 7  # 1-based line of the NOTE above


def test_redact_action_matching_is_case_insensitive(write_ged, tmp_path):
    out = str(tmp_path / "out.ged")
    stats = apply_redactions(write_ged(NOTE_GED), [_row(NOTE_LINE, action="Redact")], out)
    assert stats["redacted"] == 1
    assert REDACTED_TEXT in open(out, encoding="utf-8").read()


def test_redact_rejects_unrecognized_action(write_ged, tmp_path):
    out = tmp_path / "out.ged"
    with pytest.raises(InvalidActionError):
        apply_redactions(write_ged(NOTE_GED), [_row(NOTE_LINE, action="redcat")], str(out))
    assert not out.exists()


def test_redact_rejects_report_from_another_file(write_ged, tmp_path):
    out = tmp_path / "out.ged"
    with pytest.raises(ReportMismatchError):
        apply_redactions(write_ged(NOTE_GED), [_row(3)], str(out))  # line 3 is GEDC
    with pytest.raises(ReportMismatchError):
        apply_redactions(write_ged(NOTE_GED), [_row(NOTE_LINE, NOTE_LINE + 1)], str(out))
    assert not out.exists()


def test_privatize_wraps_only_matched_phrase(write_ged, tmp_path):
    out = str(tmp_path / "out.ged")
    rows = [
        _row(NOTE_LINE, action="privatize", term="alcoholism"),
        _row(NOTE_LINE, action="keep", term="drug problem"),
    ]
    stats = apply_redactions(write_ged(NOTE_GED), rows, out)
    content = open(out, encoding="utf-8").read()
    assert stats["privatized"] == 1
    assert "1 NOTE He struggled with [[alcoholism]] and a drug problem too." in content


def test_redact_wins_over_privatize_on_same_note(write_ged, tmp_path):
    out = str(tmp_path / "out.ged")
    rows = [
        _row(NOTE_LINE, action="privatize", term="alcoholism"),
        _row(NOTE_LINE, action="redact", term="drug problem"),
    ]
    stats = apply_redactions(write_ged(NOTE_GED), rows, out)
    content = open(out, encoding="utf-8").read()
    assert (stats["redacted"], stats["privatized"]) == (1, 0)
    assert "alcoholism" not in content


def test_privatize_expands_word_stem_to_whole_word(write_ged, tmp_path):
    out = str(tmp_path / "out.ged")
    apply_redactions(
        write_ged(NOTE_GED), [_row(NOTE_LINE, action="privatize", term="alcoholi")], out
    )
    content = open(out, encoding="utf-8").read()
    assert "[[alcoholism]]" in content
    assert "]]sm" not in content


def test_privatize_preserves_cont_line_structure(write_ged, tmp_path):
    path = write_ged(
        ged(
            "0 @I1@ INDI",
            "1 NAME Jane /Doe/",
            "1 NOTE He struggled with alcoholism for years.",
            "2 CONT He was never treated for it.",
        )
    )
    out = str(tmp_path / "out.ged")
    apply_redactions(path, [_row(7, 8, action="privatize", term="alcoholism")], out)
    lines = open(out, encoding="utf-8").read().splitlines()
    assert lines[6] == "1 NOTE He struggled with [[alcoholism]] for years."
    assert lines[7] == "2 CONT He was never treated for it."


def test_redacts_nested_and_shared_notes_at_their_own_level(write_ged, tmp_path):
    path = write_ged(EVERYWHERE)
    rows = [
        r
        for r in _rows(path, tmp_path=tmp_path)
        if r["found_in"] in ("Shared note", "Person > Death > Note")
    ]
    for r in rows:
        r["action"] = "redact"
    out = str(tmp_path / "out.ged")
    apply_redactions(path, rows, out)
    lines = open(out, encoding="utf-8").read().splitlines()
    assert f"2 NOTE {REDACTED_TEXT}" in lines
    assert f"0 @N1@ NOTE {REDACTED_TEXT}" in lines
    assert not any("prison" in ln for ln in lines)  # the CONC line went too
    assert read_gedcom(out).individuals[0].name == "Jane Doe"  # output still parses


def test_rewrite_wraps_long_and_multiline_text(write_ged, tmp_path):
    long_text = "word " * 100 + "\nsecond paragraph"
    out = str(tmp_path / "out.ged")
    apply_redactions(write_ged(NOTE_GED), [_row(NOTE_LINE, action=f"rewrite:{long_text}")], out)
    lines = open(out, encoding="utf-8").read().splitlines()
    assert all(len(ln) <= 255 for ln in lines)
    assert any(ln.startswith("2 CONC ") for ln in lines)
    assert "2 CONT second paragraph" in lines
    note = read_gedcom(out).individuals[0].notes[0].text
    assert note == long_text.strip()


def test_untouched_lines_are_copied_exactly(write_ged, tmp_path):
    path = write_ged(EVERYWHERE, newline="\r\n")
    out = str(tmp_path / "out.ged")
    apply_redactions(path, [], out)
    assert open(out, "rb").read() == open(path, "rb").read()


def test_privatize_phrase_split_across_lines(write_ged, tmp_path):
    path = write_ged(ged("0 @I1@ INDI", "1 NOTE There was a divorce", "2 CONT scandal in town."))
    rows = _rows(path, tmp_path=tmp_path)
    assert [r["matched_term"] for r in rows] == ["divorce scandal"]
    rows[0]["action"] = "privatize"
    out = str(tmp_path / "out.ged")
    apply_redactions(path, rows, out)
    assert (
        read_gedcom(out).individuals[0].notes[0].text == "There was a [[divorce\nscandal]] in town."
    )

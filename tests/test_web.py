"""The JSON entry points the browser app calls (gedcom_privacy_guard/web.py)."""

import json

from conftest import SAMPLE
from openpyxl import load_workbook

from gedcom_privacy_guard.redactor import REDACTED_TEXT
from gedcom_privacy_guard.web import redact_from_json, scan_to_json


def test_scan_returns_rows_and_writes_report(tmp_path):
    report = tmp_path / "review.xlsx"
    result = json.loads(scan_to_json(SAMPLE, str(report)))
    assert result["people"] == 4
    assert result["rows"], "sample.ged has flagged notes"
    first = result["rows"][0]
    assert {"name", "line", "line_end", "snippet", "action"} <= first.keys()
    assert first["matched_term"].lower() in result["texts"][str(first["line"])].lower()
    ws = load_workbook(report)["Review"]
    assert ws.max_row == len(result["rows"]) + 1


def test_scan_reports_progress_from_start_to_finish(tmp_path):
    calls = []
    scan_to_json(SAMPLE, str(tmp_path / "review.xlsx"), lambda f, stage: calls.append((f, stage)))
    fractions = [f for f, _ in calls]
    assert fractions[0] == 0 and fractions[-1] == 1
    assert fractions == sorted(fractions)
    assert all(stage for _, stage in calls)


def test_redact_applies_actions_from_the_page(tmp_path):
    rows = json.loads(scan_to_json(SAMPLE, str(tmp_path / "review.xlsx")))["rows"]
    for row in rows:
        row["action"] = "redact"
    out = tmp_path / "cleaned.ged"
    stats = json.loads(redact_from_json(SAMPLE, json.dumps(rows), str(out)))
    assert stats["redacted"] >= 1
    assert REDACTED_TEXT in out.read_text(encoding="utf-8")


def test_blank_actions_keep_everything(tmp_path):
    rows = json.loads(scan_to_json(SAMPLE, str(tmp_path / "review.xlsx")))["rows"]
    out = tmp_path / "cleaned.ged"
    stats = json.loads(redact_from_json(SAMPLE, json.dumps(rows), str(out)))
    assert stats["redacted"] == stats["rewritten"] == stats["privatized"] == 0


def test_facts_describe_the_file_without_its_contents(write_ged):
    path = write_ged(
        "\n".join(
            [
                "0 HEAD",
                "1 SOUR Legacy",
                "2 NAME Legacy (R)",
                "2 VERS 10.0",
                "2 CORP Millennia",
                "3 ADDR 12 Private Lane, Hometown",
                "1 FILE Cornelison Family Secrets.ged",
                "1 SUBM @U1@",
                "1 GEDC",
                "2 VERS 5.5.1",
                "1 CHAR ANSI",
                "0 @U1@ SUBM",
                "1 NAME Jane Submitter",
                "0 @I1@ INDI",
                "1 NAME Ann /Smith/",
                "1 _UID 1234",
                "1 NOTE Her son is in prison.",
                "2 CONT 0 Tuesday was hard",  # a CONT whose text looks like a record
                "0 Smithfield notes",
                "0 TRLR",
            ]
        )
        + "\n",
        encoding="cp1252",
        newline="\r\n",
    )
    from gedcom_privacy_guard.parser import read_gedcom
    from gedcom_privacy_guard.web import file_facts

    facts = file_facts(read_gedcom(path))
    assert facts["program"] == "Legacy"
    assert facts["program_name"] == "Legacy (R)"
    assert facts["program_version"] == "10.0"
    assert facts["gedcom_version"] == "5.5.1"
    assert facts["declared_charset"] == "ANSI"
    assert facts["line_ending"] == "CRLF"
    assert facts["records"] == {"SUBM": 1, "INDI": 1, "other": 1}
    assert facts["custom_tags"] == {"_UID": 1}
    text = json.dumps(facts)
    for private in (
        "Private Lane",
        "Secrets",
        "Submitter",
        "Smith",
        "prison",
        "Tuesday",
        "Millennia",
    ):
        assert private not in text


def test_scan_includes_the_facts(tmp_path):
    result = json.loads(scan_to_json(SAMPLE, str(tmp_path / "review.xlsx")))
    assert result["facts"]["program"] == "GedcomPrivacyGuardSample"
    assert result["facts"]["records"]["INDI"] == 4
